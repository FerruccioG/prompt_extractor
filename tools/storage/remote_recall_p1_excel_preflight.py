#!/usr/bin/env python3
import json
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
BOOK = Path("/mnt/remote/Remote_Opportunities_Master_List_Expanded_Sources.xlsx")
SHEET = "Source Directory"
EXPANDED = ROOT / "data/remote/source_records_recall_p1_expanded.jsonl"
PROMOTIONS = ROOT / "data/remote/recall_p1_promote_now_refreshed.jsonl"

ALIASES = {
    "gcstechtalent.com": "gcsrecruitment.com",
    "javascript.jobs": "jsremotely.com",
    "diversityjobs.com": "latpro.com",
}

def rows(path):
    out=[]
    with path.open(encoding="utf-8") as f:
        for line in f:
            line=line.strip()
            if line:
                out.append(json.loads(line))
    return out

def host(v):
    s=str(v or "").strip().lower()
    if not s:
        return ""
    p=urlparse(s if "://" in s else "https://"+s)
    h=(p.hostname or "").strip(".")
    if h.startswith("www."):
        h=h[4:]
    return ALIASES.get(h,h)

def main():
    expanded=rows(EXPANDED)
    promos=rows(PROMOTIONS)
    promo_sources={str(r.get("Source","")).lower() for r in promos}

    wb=load_workbook(BOOK, read_only=True, data_only=False)
    ws=wb[SHEET]

    header_row=3
    headers={}
    for c in range(1, ws.max_column+1):
        v=ws.cell(header_row,c).value
        if v:
            headers[str(v).strip().casefold()]=c

    src_col=headers["source"]
    url_col=headers["url"]
    score_col=headers.get("candidate ↔ source relationship score")
    tier_col=headers.get("relationship tier")

    by_source={}
    by_host={}
    for r in range(header_row+1, ws.max_row+1):
        src=str(ws.cell(r,src_col).value or "").strip().lower()
        url=str(ws.cell(r,url_col).value or "").strip()
        if src and src not in by_source:
            by_source[src]=r
        h=host(url)
        if h and h not in by_host:
            by_host[h]=r

    details=[]
    used={}
    for rec in expanded:
        src=str(rec.get("Source","")).strip().lower()
        rh=host(rec.get("URL") or src)
        r=by_source.get(src) or by_host.get(rh)
        d={"source":src,"row":r,"action":"update" if r else "append","promotion":src in promo_sources}
        if r:
            d["existing_source"]=str(ws.cell(r,src_col).value or "")
            d["existing_url"]=str(ws.cell(r,url_col).value or "")
            d["existing_score"]=ws.cell(r,score_col).value if score_col else None
            d["existing_tier"]=ws.cell(r,tier_col).value if tier_col else None
            used.setdefault(r,[]).append(src)
        details.append(d)

    dup={str(k):v for k,v in used.items() if len(v)>1}
    promo=[d for d in details if d["promotion"]]
    summary={
        "status":"ok" if not dup else "error",
        "expanded_records":len(expanded),
        "promotion_records":len(promos),
        "would_update":sum(1 for d in details if d["action"]=="update"),
        "would_append":sum(1 for d in details if d["action"]=="append"),
        "duplicate_target_rows":dup,
        "promotion_existing_legacy_rows":sum(1 for d in promo if d["action"]=="update" and not d.get("existing_score")),
        "promotion_existing_golden_rows":sum(1 for d in promo if d["action"]=="update" and d.get("existing_score")),
        "promotion_physical_appends":sum(1 for d in promo if d["action"]=="append"),
        "promotion_details":promo,
    }
    out=ROOT/"data/remote/recall_p1_excel_preflight.json"
    out.write_text(json.dumps(summary,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps(summary,ensure_ascii=False))
    wb.close()

if __name__=="__main__":
    main()
