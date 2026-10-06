#!/usr/bin/env python3
"""
excel_source_integrity_validator.py

Read-only post-write integrity validator for the Golden Remote Opportunities
workbook.

Validates that:
- workbook and Source Directory are readable;
- the 21 Golden relationship columns exist;
- every source in data/remote/source_records.jsonl is present exactly once;
- canonical host/source matching does not reveal duplicate Golden rows;
- key multidimensional scores and tier values match the JSONL source records;
- Tier distribution of the 38 published records is preserved.

This script never writes to the workbook.
"""

from __future__ import annotations

import json
import os
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook

ROOT_DIR = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = Path("/mnt/remote/Remote_Opportunities_Master_List_Expanded_Sources.xlsx")
DEFAULT_INPUT = ROOT_DIR / "data/remote/source_records.jsonl"
DEFAULT_SHEET = "Source Directory"

GOLDEN_HEADERS = [
    "Candidate ↔ Source Relationship Score",
    "Relationship Tier",
    "Candidate Fit",
    "Remote Strength",
    "Ireland/EU Accessibility",
    "Opportunity Density",
    "Specialization Fit",
    "Hidden Potential",
    "Directness",
    "Trust/Risk",
    "Earning Potential",
    "Signal Quality",
    "Freshness/Activity",
    "Generic Relationship Score",
    "Verification Status",
    "Email Evidence Count",
    "Discovery Evidence Count",
    "Direct Evidence URL Count",
    "Profile Signal Groups",
    "Representative Evidence URLs",
    "Scoring Notes",
]

KEY_COMPARE_FIELDS = [
    "Candidate ↔ Source Relationship Score",
    "Relationship Tier",
    "Candidate Fit",
    "Remote Strength",
    "Ireland/EU Accessibility",
    "Opportunity Density",
    "Specialization Fit",
    "Hidden Potential",
    "Directness",
    "Trust/Risk",
    "Candidate Cost",
    "Earning Potential",
    "Signal Quality",
    "Freshness/Activity",
    "Generic Relationship Score",
    "Verification Status",
    "Email Evidence Count",
    "Discovery Evidence Count",
    "Direct Evidence URL Count",
    "Profile Signal Groups",
]

HOST_ALIASES = {
    "m.hays.ie": "hays.ie",
    "onlineapi-internet.hays.com": "hays.ie",
    "candidate-support.crossover.com": "crossover.com",
    "apply2.computerfutures.com": "computerfutures.com",
    "ie.gcsrecruitment.com": "gcsrecruitment.com",
    "gcstechtalent.com": "gcsrecruitment.com",
    "in.indeed.com": "indeed.com",
    "uk.indeed.com": "indeed.com",
    "in.talent.com": "talent.com",
    "ie.neuvoo.com": "talent.com",
    "jobs.zendesk.com": "zendesk.com",
    "careers.microsoft.com": "microsoft.com",
    "members.microsoft.com": "microsoft.com",
    "careers.bunq.com": "bunq.com",
    "press.bunq.com": "bunq.com",
    "careers.accelerationpartners.com": "accelerationpartners.com",
    "remotelyx-3.careers-page.com": "remotelyx.com",
    "javascript.jobs": "jsremotely.com",
    "diversityjobs.com": "latpro.com",
}


def clean(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def key(value: Any) -> str:
    return clean(value).casefold()


def normalized_host(value: Any) -> str:
    raw = clean(value)
    if not raw:
        return ""
    try:
        parsed = urlparse(raw if "://" in raw else f"https://{raw}")
        host = (parsed.hostname or "").casefold().strip(".")
    except Exception:
        return ""
    if host.startswith("www."):
        host = host[4:]
    return HOST_ALIASES.get(host, host)


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as fh:
        for line_no, raw in enumerate(fh, start=1):
            line = raw.strip()
            if not line:
                continue
            try:
                value = json.loads(line)
            except json.JSONDecodeError as exc:
                raise RuntimeError(f"Invalid JSON at {path}:{line_no}: {exc}") from exc
            if not isinstance(value, dict):
                raise RuntimeError(f"Expected JSON object at {path}:{line_no}")
            rows.append(value)
    return rows


def find_header_row(ws, max_scan_rows: int = 15) -> int:
    for row in range(1, min(ws.max_row, max_scan_rows) + 1):
        vals = {
            key(ws.cell(row=row, column=col).value)
            for col in range(1, ws.max_column + 1)
        }
        if {"source", "url"}.issubset(vals):
            return row
    raise RuntimeError("Could not locate Source/URL header row.")


def header_map(ws, header_row: int) -> dict[str, int]:
    out: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = clean(ws.cell(row=header_row, column=col).value)
        if value:
            out[value.casefold()] = col
    return out


def row_value(ws, row: int, headers: dict[str, int], name: str) -> Any:
    col = headers.get(name.casefold())
    return ws.cell(row=row, column=col).value if col else None


def comparable(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def main() -> int:
    workbook = Path(os.getenv("REMOTE_EXCEL_OUTPUT_PATH", str(DEFAULT_WORKBOOK)))
    source_input = Path(os.getenv("REMOTE_SOURCE_RECORDS_PATH", str(DEFAULT_INPUT)))
    sheet = os.getenv("REMOTE_EXCEL_SOURCE_SHEET", DEFAULT_SHEET).strip() or DEFAULT_SHEET

    if not workbook.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook}")
    if not source_input.exists():
        raise FileNotFoundError(f"Source records not found: {source_input}")

    records = load_jsonl(source_input)
    wb = load_workbook(workbook, data_only=False, read_only=False)
    try:
        if sheet not in wb.sheetnames:
            raise RuntimeError(f"Worksheet not found: {sheet}")

        ws = wb[sheet]
        header_row = find_header_row(ws)
        headers = header_map(ws, header_row)

        missing_headers = [h for h in GOLDEN_HEADERS if h.casefold() not in headers]

        source_index: dict[str, list[int]] = defaultdict(list)
        host_index: dict[str, list[int]] = defaultdict(list)

        for row in range(header_row + 1, ws.max_row + 1):
            source = key(row_value(ws, row, headers, "Source"))
            url = row_value(ws, row, headers, "URL")
            host = normalized_host(url)
            if source:
                source_index[source].append(row)
            if host:
                host_index[host].append(row)

        missing_sources: list[str] = []
        duplicate_sources: dict[str, list[int]] = {}
        mismatches: list[dict[str, Any]] = []
        matched_rows: dict[str, int] = {}
        published_tiers = Counter()

        for record in records:
            source = key(record.get("Source"))
            url = record.get("URL")
            host = normalized_host(url)

            candidate_rows = source_index.get(source, [])
            if not candidate_rows and host:
                candidate_rows = host_index.get(host, [])

            if not candidate_rows:
                missing_sources.append(source)
                continue

            if len(candidate_rows) > 1:
                duplicate_sources[source] = candidate_rows
                continue

            row = candidate_rows[0]
            matched_rows[source] = row
            published_tiers[str(record.get("Relationship Tier", ""))] += 1

            for field in KEY_COMPARE_FIELDS:
                expected = comparable(record.get(field))
                actual = comparable(row_value(ws, row, headers, field))
                if expected != actual:
                    mismatches.append({
                        "source": source,
                        "row": row,
                        "field": field,
                        "expected": expected,
                        "actual": actual,
                    })

        # Duplicate canonical hosts among the 38 published relationships only.
        published_host_rows: dict[str, list[tuple[str, int]]] = defaultdict(list)
        for record in records:
            source = key(record.get("Source"))
            row = matched_rows.get(source)
            host = normalized_host(record.get("URL"))
            if row and host:
                published_host_rows[host].append((source, row))

        duplicate_published_hosts = {
            host: vals
            for host, vals in published_host_rows.items()
            if len(vals) > 1
        }

        expected_tiers = Counter(str(r.get("Relationship Tier", "")) for r in records)

        status = "ok"
        if missing_headers or missing_sources or duplicate_sources or mismatches or duplicate_published_hosts:
            status = "failed"
        if published_tiers != expected_tiers:
            status = "failed"

        result = {
            "status": status,
            "workbook": str(workbook),
            "worksheet": sheet,
            "header_row": header_row,
            "mapped_columns": len(headers),
            "worksheet_rows_after_header": max(0, ws.max_row - header_row),
            "source_records_expected": len(records),
            "source_records_matched": len(matched_rows),
            "golden_headers_expected": len(GOLDEN_HEADERS),
            "golden_headers_missing": missing_headers,
            "tier_distribution_expected": dict(expected_tiers),
            "tier_distribution_matched": dict(published_tiers),
            "missing_sources": missing_sources,
            "duplicate_source_rows": duplicate_sources,
            "duplicate_published_canonical_hosts": duplicate_published_hosts,
            "field_mismatch_count": len(mismatches),
            "field_mismatches": mismatches[:100],
        }

        print(json.dumps(result, indent=2, ensure_ascii=False))
        return 0 if status == "ok" else 1
    finally:
        wb.close()


if __name__ == "__main__":
    raise SystemExit(main())
