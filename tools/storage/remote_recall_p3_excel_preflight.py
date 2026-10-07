#!/usr/bin/env python3
import json
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
BOOK=Path("/mnt/remote/Remote_Opportunities_Master_List_Expanded_Sources.xlsx")
EXPANDED=ROOT/"data/remote/source_records_recall_p3_expanded.jsonl"
PROMOTIONS=ROOT/"data/remote/recall_p3_promote_now_refreshed.jsonl"
OUT=ROOT/"data/remote/recall_p3_excel_preflight.json"
ALIASES={"gcstechtalent.com":"gcsrecruitment.com","javascript.jobs":"jsremotely.com","diversityjobs.com":"latpro.com","angel.co":"wellfound.com","remotive.io":"remotive.com","robertwalters.co.uk":"robertwalters.com"}

def rows(p):
    with p.open(encoding="utf-8") as f:
        return [json.loads(x) for x in f if x.strip()]

def host(v):
    s=str(v or "").strip().lower()
    if not s:return ""
    p=urlparse(s if "://" in s else "https://"+s)
    h=(p.hostname or "").strip(".")
    if h.startswith("www."):h=h[4:]
    return ALIASES.get(h,h)

expanded=rows(EXPANDED); promos=rows(PROMOTIONS)
promo_sources={str(r.get("Source","")).lower() for r in promos}
wb=load_workbook(BOOK,read_only=True,data_only=False); ws=wb["Source Directory"]
headers={str(ws.cell(3,c).value).strip().casefold():c for c in range(1,ws.max_column+1) if ws.cell(3,c).value}
src_col=headers["source"]; url_col=headers["url"]
score_col=headers.get("candidate ↔ source relationship score"); tier_col=headers.get("relationship tier")
by_source={}; by_host={}
for r in range(4,ws.max_row+1):
    src=str(ws.cell(r,src_col).value or "").strip().lower()
    u=str(ws.cell(r,url_col).value or "").strip()
    if src and src not in by_source:by_source[src]=r
    h=host(u)
    if h and h not in by_host:by_host[h]=r
details=[]; used={}
for rec in expanded:
    src=str(rec.get("Source","")).strip().lower()
    r=by_source.get(src) or by_host.get(host(rec.get("URL") or src))
    d={"source":src,"row":r,"action":"update" if r else "append","promotion":src in promo_sources}
    if r:
        d["existing_score"]=ws.cell(r,score_col).value if score_col else None
        d["existing_tier"]=ws.cell(r,tier_col).value if tier_col else None
        used.setdefault(r,[]).append(src)
    details.append(d)
dup={str(k):v for k,v in used.items() if len(v)>1}
promo=[d for d in details if d["promotion"]]
summary={"status":"ok" if not dup else "error","expanded_records":len(expanded),"promotion_records":len(promos),"would_update":sum(d["action"]=="update" for d in details),"would_append":sum(d["action"]=="append" for d in details),"duplicate_target_rows":dup,"promotion_existing_legacy_rows":sum(d["action"]=="update" and not d.get("existing_score") for d in promo),"promotion_existing_golden_rows":sum(d["action"]=="update" and bool(d.get("existing_score")) for d in promo),"promotion_physical_appends":sum(d["action"]=="append" for d in promo),"promotion_details":promo}
OUT.write_text(json.dumps(summary,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
print(json.dumps(summary,ensure_ascii=False)); wb.close()
